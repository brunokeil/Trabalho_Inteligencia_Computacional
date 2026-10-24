#!/usr/bin/env python3
"""Padronização, conversão e renomeação das instâncias em instances/.
Preserva os 52 arquivos originais em instances/originais/
Gera 41 instâncias únicas, limpas e validadas no formato padrão DCC136.
Gera instances/README.md e instances/manifesto.csv para rastreabilidade completa.
"""

import os, sys, shutil, hashlib, re, csv
from decimal import Decimal
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
INSTANCES_DIR = ROOT / 'instances'
ORIG_DIR = INSTANCES_DIR / 'originais'

GROUPS = [
    # (novo_nome, arquivo_fonte_primario, descricao, [arquivos_equivalentes], tipo_conversao)
    ('instancia_01_prof_A.txt', 'instancia_A.txt', 'Professora Luciana - Referência A', ['instancia_A.txt', 'instancia_A 2.txt'], 'standard'),
    ('instancia_02_prof_B.txt', 'instancia_B.txt', 'Professora Luciana - Referência B', ['instancia_B.txt', 'instancia_B 2.txt'], 'standard'),
    ('instancia_03_pequena_100.txt', 'instancia1 2.txt', 'Modelo Pequeno (n=100)', ['instancia1 2.txt'], 'standard'),
    ('instancia_04_pequena_109.txt', 'instancia2 2.txt', 'Modelo Pequeno (n=109)', ['instancia2 2.txt'], 'standard'),
    ('instancia_05_turma_2550_1.txt', 'instancia1.txt', 'Turma DCC136 - Instância 1 (n=2550)', ['instancia1.txt'], 'standard'),
    ('instancia_06_turma_2550_2.txt', 'instancia2.txt', 'Turma DCC136 - Instância 2 (n=2550)', ['instancia2.txt'], 'standard'),
    ('instancia_07_malvacini_1.txt', 'instancia_1 Arthur Malvacini.txt', 'Arthur Malvacini - Instância 1', ['instancia_1 Arthur Malvacini.txt', 'instancia_1 (1).txt'], 'standard'),
    ('instancia_08_malvacini_2.txt', 'instancia_2 Arthur Malvacini.txt', 'Arthur Malvacini - Instância 2', ['instancia_2 Arthur Malvacini.txt', 'instancia_2 (1).txt'], 'standard'),
    ('instancia_09_gustavo_isabela_base.txt', 'instancia_base_Gustavo_Isabela.txt', 'Gustavo & Isabela - Base', ['instancia_base_Gustavo_Isabela.txt', 'instancia_base.txt', 'instancia_base(1).txt'], 'standard'),
    ('instancia_10_gustavo_isabela_dif.txt', 'instancia_diferenciada_Gustavo_Isabela.txt', 'Gustavo & Isabela - Diferenciada', ['instancia_diferenciada_Gustavo_Isabela.txt', 'instancia_diferenciada.txt', 'instancia_diferenciada(1).txt'], 'standard'),
    ('instancia_11_felipe_mickaell_1.txt', 'instancia_felipe_mickaell.txt', 'Felipe & Mickaell - Instância 1', ['instancia_felipe_mickaell.txt'], 'standard'),
    ('instancia_12_felipe_mickaell_2.txt', 'instancia_mickaell_felipe.txt', 'Felipe & Mickaell - Instância 2', ['instancia_mickaell_felipe.txt'], 'standard'),
    ('instancia_13_g01_360.txt', 'instancia_1.txt', 'Grupo 01 - Instância 1 (n=360)', ['instancia_1.txt'], 'standard'),
    ('instancia_14_g01_382.txt', 'instancia_2.txt', 'Grupo 01 - Instância 2 (n=382)', ['instancia_2.txt'], 'standard'),
    ('instancia_15_g02_370.txt', 'instancia_1 (1)(1).txt', 'Grupo 02 - Instância 1 (n=370)', ['instancia_1 (1)(1).txt'], 'standard'),
    ('instancia_16_g02_362.txt', 'instancia_2 (1)(1).txt', 'Grupo 02 - Instância 2 (n=362)', ['instancia_2 (1)(1).txt'], 'standard'),
    ('instancia_17_g03_3452.txt', 'instancia_1(1).txt', 'Grupo 03 - Instância 1 (n=3452)', ['instancia_1(1).txt'], 'standard'),
    ('instancia_18_g03_4454.txt', 'instancia_2(1).txt', 'Grupo 03 - Instância 2 (n=4454)', ['instancia_2(1).txt'], 'standard'),
    ('instancia_19_g04_4213.txt', 'instancia_1(2).txt', 'Grupo 04 - Instância 1 (n=4213)', ['instancia_1(2).txt'], 'standard'),
    ('instancia_20_g04_316.txt', 'instancia_2(2).txt', 'Grupo 04 - Instância 2 (n=316)', ['instancia_2(2).txt'], 'standard'),
    ('instancia_21_g05_316.txt', 'instancia_1(3).txt', 'Grupo 05 - Instância 1 (n=316)', ['instancia_1(3).txt'], 'standard'),
    ('instancia_22_g05_3900.txt', 'instancia_2(3).txt', 'Grupo 05 - Instância 2 (n=3900)', ['instancia_2(3).txt'], 'standard'),
    ('instancia_23_g06_4656.txt', 'instancia_1(4).txt', 'Grupo 06 - Instância 1 (n=4656)', ['instancia_1(4).txt', 'instancia_1(6).txt'], 'variant_4'),
    ('instancia_24_g06_7038.txt', 'instancia_2(4).txt', 'Grupo 06 - Instância 2 (n=7038)', ['instancia_2(4).txt', 'instancia_2(6).txt'], 'variant_4'),
    ('instancia_25_g07_5917.txt', 'instancia_1(5).txt', 'Grupo 07 - Instância 1 (n=5917)', ['instancia_1(5).txt'], 'standard'),
    ('instancia_26_g07_3858.txt', 'instancia_2(5).txt', 'Grupo 07 - Instância 2 (n=3858)', ['instancia_2(5).txt'], 'standard'),
    ('instancia_27_g08_4175.txt', 'instancia_1(7).txt', 'Grupo 08 - Instância 1 (n=4175)', ['instancia_1(7).txt'], 'variant_7'),
    ('instancia_28_g08_4784.txt', 'instancia_2(7).txt', 'Grupo 08 - Instância 2 (n=4784)', ['instancia_2(7).txt'], 'variant_7'),
    ('instancia_29_g09_3531.txt', 'instancia_1(8).txt', 'Grupo 09 - Instância 1 (n=3531)', ['instancia_1(8).txt'], 'variant_8'),
    ('instancia_30_g09_4131.txt', 'instancia_2(8).txt', 'Grupo 09 - Instância 2 (n=4131)', ['instancia_2(8).txt'], 'variant_8'),
    ('instancia_31_g10_base.txt', 'instancia_base(2).txt', 'Grupo 10 - Base (n=4523)', ['instancia_base(2).txt'], 'variant_base'),
    ('instancia_32_g10_dif.txt', 'instancia_diferenciada(2).txt', 'Grupo 10 - Diferenciada (n=4378)', ['instancia_diferenciada(2).txt'], 'variant_base'),
    ('instancia_33_g11_base.txt', 'instancia_base(3).txt', 'Grupo 11 - Base (n=4543)', ['instancia_base(3).txt'], 'variant_base'),
    ('instancia_34_g11_dif.txt', 'instancia_diferenciada(3).txt', 'Grupo 11 - Diferenciada (n=4254)', ['instancia_diferenciada(3).txt'], 'variant_base'),
    ('instancia_35_g12_formatada_1.txt', 'instancia_1_f.txt', 'Grupo 12 - Formatada 1 (n=3860)', ['instancia_1_f.txt'], 'variant_f'),
    ('instancia_36_g12_formatada_2.txt', 'instancia_2_f.txt', 'Grupo 12 - Formatada 2 (n=3860)', ['instancia_2_f.txt'], 'variant_f'),
    ('instancia_37_g13_3696.txt', 'instancias1.txt', 'Grupo 13 - Instância 1 (n=3696)', ['instancias1.txt'], 'standard'),
    ('instancia_38_g13_4124.txt', 'instancias2.txt', 'Grupo 13 - Instância 2 (n=4124)', ['instancias2.txt'], 'standard'),
    ('instancia_39_g14_excel_1.txt', 'instancia_1.xlsx', 'Grupo 14 - Excel 1 (n=4213)', ['instancia_1.xlsx'], 'xlsx'),
    ('instancia_40_g14_excel_2.txt', 'instancia_2.xlsx', 'Grupo 14 - Excel 2 (n=3918)', ['instancia_2.xlsx'], 'xlsx'),
    ('instancia_41_grande_5287.txt', 'instancia_2_grande.txt', 'Instância Grande Extra (n=5287)', ['instancia_2_grande.txt', 'instancia_3_extra.txt'], 'standard'),
]

def convert_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws_res = wb['Resumo']
    n = None; m = None
    for r in ws_res.iter_rows(values_only=True):
        if r[0] and 'candidatos' in str(r[0]).lower(): n = int(r[1])
        if r[0] and 'escolas' in str(r[0]).lower(): m = int(r[1])
    cands = []
    ws_c = wb['Candidatos']
    for i, r in enumerate(ws_c.iter_rows(values_only=True)):
        if i == 0: continue
        cid, cep, tipo = str(r[0]), str(r[1]).replace('-', ''), str(r[4])
        cands.append(f'{cid}, {cep}, {tipo}, -21.7600, -43.3500')
    salas_by_esc = {}
    ws_s = wb['Salas']
    for i, r in enumerate(ws_s.iter_rows(values_only=True)):
        if i == 0: continue
        eid, cap = str(r[0]), int(r[2])
        salas_by_esc.setdefault(eid, []).append(str(cap))
    schools = []
    ws_e = wb['Escolas']
    for i, r in enumerate(ws_e.iter_rows(values_only=True)):
        if i == 0: continue
        eid, cep, nr = str(r[0]), str(r[1]).replace('-', ''), int(r[4])
        caps = ';'.join(salas_by_esc[eid])
        schools.append(f'{eid}, {cep}, -21.7600, -43.3500, {nr}, {caps}')
    matrix = []
    ws_m = wb['Matriz_Distancias']
    for i, r in enumerate(ws_m.iter_rows(values_only=True)):
        if i == 0: continue
        dists = [f'{Decimal(str(d)):.4f}' for d in r[1:m+1]]
        matrix.append(', '.join(dists))
    return '\n'.join([f'NUM_CANDIDATOS: {n}', f'NUM_ESCOLAS: {m}', '# CANDIDATOS'] + cands + ['# ESCOLAS'] + schools + ['# MATRIZ_DISTANCIAS'] + matrix)

def convert_variant_base(lines):
    p = 0
    def nxt():
        nonlocal p; l = lines[p]; p += 1; return l
    h1 = nxt().split(':'); n = int(h1[1].strip())
    h2 = nxt().split(':'); m = int(h2[1].strip())
    cands = []
    for _ in range(n):
        parts = [x.strip() for x in nxt().split(',')]
        cands.append(f'{parts[0]}, {parts[1]}, {parts[2]}, -21.7600, -43.3500')
    schools = []
    for _ in range(m):
        parts = [x.strip() for x in nxt().split(',')]
        eid, cep, nr, caps = parts[0], parts[1], parts[4], parts[5]
        schools.append(f'{eid}, {cep}, -21.7600, -43.3500, {nr}, {caps}')
    matrix = [nxt() for _ in range(n)]
    return '\n'.join([f'NUM_CANDIDATOS: {n}', f'NUM_ESCOLAS: {m}', '# CANDIDATOS'] + cands + ['# ESCOLAS'] + schools + ['# MATRIZ_DISTANCIAS'] + matrix)

def convert_variant_4(lines):
    n = None; m = None
    cands = []; schools = []; salas = {}; matrix = []
    mode = None
    for l in lines:
        if l.startswith('NUM_CANDIDATOS'): n = int(l.split(':')[1].strip())
        elif l.startswith('NUM_ESCOLAS'): m = int(l.split(':')[1].strip())
        elif l.startswith('# CANDIDATOS'): mode = 'cand'
        elif l.startswith('# ESCOLAS'): mode = 'esc'
        elif l.startswith('# SALAS'): mode = 'salas'
        elif l.startswith('# MATRIZ'): mode = 'mat'
        elif l.startswith('#'): continue
        elif mode == 'cand':
            parts = l.split()
            cands.append(f'{parts[0]}, {parts[1]}, {parts[4]}, -21.7600, -43.3500')
        elif mode == 'esc':
            parts = l.split()
            schools.append((parts[0], parts[1], parts[6]))
        elif mode == 'salas':
            parts = l.split()
            salas[parts[0]] = ';'.join(parts[1:])
        elif mode == 'mat':
            matrix.append(', '.join([f'{Decimal(d):.4f}' for d in l.split()]))
    esc_lines = [f'{eid}, {cep}, -21.7600, -43.3500, {nr}, {salas[eid]}' for eid, cep, nr in schools]
    return '\n'.join([f'NUM_CANDIDATOS: {n}', f'NUM_ESCOLAS: {m}', '# CANDIDATOS'] + cands + ['# ESCOLAS'] + esc_lines + ['# MATRIZ_DISTANCIAS'] + matrix)

def convert_variant_7(lines):
    p = 0
    def nxt():
        nonlocal p; l = lines[p]; p += 1; return l
    m = int(nxt())
    schools = []
    for _ in range(m):
        parts = nxt().split()
        caps_str = ';'.join(parts[4:])
        schools.append(f'E{parts[0]}, {parts[1]}, -21.7600, -43.3500, {parts[3]}, {caps_str}')
    n = int(nxt())
    cands = []
    for _ in range(n):
        parts = nxt().split()
        cands.append(f'C{parts[0]}, {parts[1]}, {parts[2]}, -21.7600, -43.3500')
    matrix = [', '.join([f'{Decimal(d):.4f}' for d in nxt().split()]) for _ in range(n)]
    return '\n'.join([f'NUM_CANDIDATOS: {n}', f'NUM_ESCOLAS: {m}', '# CANDIDATOS'] + cands + ['# ESCOLAS'] + schools + ['# MATRIZ_DISTANCIAS'] + matrix)

def convert_variant_8(lines):
    n = int(lines[0]); m = int(lines[1])
    cands = []; schools = []; salas = {}; matrix = []
    mode = None
    for l in lines[2:]:
        if l.startswith('# CANDIDATOS'): mode = 'cand'
        elif l.startswith('# ESCOLAS'): mode = 'esc'
        elif l.startswith('# SALAS'): mode = 'salas'
        elif l.startswith('# MATRIZ'): mode = 'mat'
        elif l.startswith('#'): continue
        elif mode == 'cand':
            parts = l.split()
            cands.append(f'{parts[0]}, {parts[1]}, {parts[4]}, -21.7600, -43.3500')
        elif mode == 'esc':
            parts = l.split()
            schools.append((parts[0], parts[1], parts[5]))
        elif mode == 'salas':
            parts = l.split()
            salas[parts[0]] = ';'.join(parts[1:])
        elif mode == 'mat':
            matrix.append(', '.join([f'{Decimal(d):.4f}' for d in l.split()]))
    esc_lines = [f'{eid}, {cep}, -21.7600, -43.3500, {nr}, {salas[eid]}' for eid, cep, nr in schools]
    return '\n'.join([f'NUM_CANDIDATOS: {n}', f'NUM_ESCOLAS: {m}', '# CANDIDATOS'] + cands + ['# ESCOLAS'] + esc_lines + ['# MATRIZ_DISTANCIAS'] + matrix)

def convert_variant_f(lines):
    n = None; m = None
    cands = []; schools = []; salas = {}; matrix = []
    mode = None
    for l in lines:
        if l.startswith('1. Quantidade de Candidatos'):
            n = int(re.search(r'Candidatos:\s*(\d+)', l).group(1))
        elif l.startswith('2. Quantidade de Escolas'):
            m = int(re.search(r'Escolas:\s*(\d+)', l).group(1))
        elif l.startswith('3. Dados de cada candidato'): mode = 'cand'
        elif l.startswith('4. Dados de cada escola'): mode = 'esc'
        elif l.startswith('5. Capacidade de cada sala'): mode = 'salas'
        elif l.startswith('6. Matriz de distancias'): mode = 'mat'
        elif mode == 'cand':
            m_cand = re.match(r'^(C\d+),\s*CEP\s*([0-9\-]+)\s*\([^)]*\),\s*(\S+)', l)
            if m_cand:
                cand_cep = m_cand.group(2).replace('-', '')
                cands.append(f'{m_cand.group(1)}, {cand_cep}, {m_cand.group(3)}, -21.7600, -43.3500')
        elif mode == 'esc':
            m_esc = re.match(r'^(E\d+),\s*CEP\s*([0-9\-]+)\s*\([^)]*\),\s*(\d+)\s*salas', l)
            if m_esc:
                schools.append((m_esc.group(1), m_esc.group(2).replace('-', ''), m_esc.group(3)))
        elif mode == 'salas':
            m_s = re.match(r'^(E\d+):\s*(\d+)\s*sala\(s\)\s*com\s*capacidade\s*para\s*(\d+)\s*alunos', l)
            if m_s:
                salas[m_s.group(1)] = ';'.join([m_s.group(3)] * int(m_s.group(2)))
        elif mode == 'mat':
            if l.startswith('ID'): continue
            matrix.append(', '.join([f'{Decimal(d):.4f}' for d in l.split()[1:]]))
    esc_lines = [f'{eid}, {cep}, -21.7600, -43.3500, {nr}, {salas[eid]}' for eid, cep, nr in schools]
    return '\n'.join([f'NUM_CANDIDATOS: {n}', f'NUM_ESCOLAS: {m}', '# CANDIDATOS'] + cands + ['# ESCOLAS'] + esc_lines + ['# MATRIZ_DISTANCIAS'] + matrix)

def clean_standard(lines):
    out = []
    for l in lines:
        sl = l.strip()
        if sl:
            out.append(sl)
    return '\n'.join(out)

def main():
    ORIG_DIR.mkdir(parents=True, exist_ok=True)
    
    # 1. Backup all original files
    print(f'1. Fazendo backup dos arquivos originais em {ORIG_DIR.relative_to(ROOT)}...')
    original_files = [f for f in os.listdir(INSTANCES_DIR) if not f.startswith('.') and f != 'originais' and (INSTANCES_DIR / f).is_file()]
    for f in original_files:
        src = INSTANCES_DIR / f
        dst = ORIG_DIR / f
        if not dst.exists():
            shutil.copy2(src, dst)
    print(f'   {len(original_files)} arquivos preservados no backup.')

    # 2. Convert and write each group to target
    print('2. Convertendo e padronizando as 41 instâncias...')
    manifest_rows = []
    
    for target_name, src_name, desc, duplicates, kind in GROUPS:
        src_path = ORIG_DIR / src_name
        if kind == 'xlsx':
            content = convert_xlsx(src_path)
        else:
            with open(src_path, 'r', encoding='utf-8', errors='ignore') as f:
                lines = f.readlines()
            if kind == 'standard':
                content = clean_standard(lines)
            elif kind == 'variant_base':
                content = convert_variant_base([l.strip() for l in lines if l.strip() and not l.strip().startswith('#')])
            elif kind == 'variant_4':
                content = convert_variant_4(lines)
            elif kind == 'variant_7':
                content = convert_variant_7(lines)
            elif kind == 'variant_8':
                content = convert_variant_8(lines)
            elif kind == 'variant_f':
                content = convert_variant_f(lines)
        
        target_path = INSTANCES_DIR / target_name
        with open(target_path, 'w', encoding='utf-8') as f:
            f.write(content + '\n')
            
        sha256 = hashlib.sha256((content + '\n').encode('utf-8')).hexdigest()
        
        # Read stats
        lines_clean = [l.strip() for l in content.splitlines() if l.strip() and not l.strip().startswith('#')]
        n = int(lines_clean[0].split(':')[1].strip())
        m = int(lines_clean[1].split(':')[1].strip())
        schools = lines_clean[2+n:2+n+m]
        caps = [sum(map(int, s.split(',')[5].split(';'))) for s in schools]
        total_cap = sum(caps)
        utilizacao = 100.0 * n / total_cap if total_cap else 0
        
        manifest_rows.append({
            'arquivo': target_name,
            'origem_primaria': src_name,
            'grupo_descricao': desc,
            'arquivos_duplicados': '; '.join(duplicates),
            'n_candidatos': n,
            'm_escolas': m,
            'capacidade_total': total_cap,
            'utilizacao_pct': f'{utilizacao:.2f}%',
            'sha256': sha256
        })
        print(f'   -> {target_name:32s} (n={n:4d}, m={m:2d}, cap={total_cap:5d}) OK')

    # 3. Clean up non-standard / old root files
    print('3. Limpando arquivos legados da raiz de instances/ (mantidos em originais/)...')
    for f in original_files:
        p = INSTANCES_DIR / f
        if p.is_file() and p.name not in [g[0] for g in GROUPS]:
            p.unlink()

    # 4. Generate manifesto.csv
    csv_path = INSTANCES_DIR / 'manifesto.csv'
    with open(csv_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=list(manifest_rows[0].keys()))
        writer.writeheader()
        writer.writerows(manifest_rows)
    print(f'   Salvo: {csv_path.relative_to(ROOT)}')

    # 5. Generate README.md
    md_lines = [
        '# Instâncias Padronizadas — DCC136 (2026.3)',
        '',
        'Este diretório contém todas as instâncias compartilhadas pela turma de Inteligência Computacional (DCC136 / UFJF / 2026.3), padronizadas para o formato canônico aceito pelo executável `./alocacao`.',
        '',
        '## Informações gerais',
        '',
        f'- **Total de instâncias únicas:** {len(manifest_rows)}',
        '- **Arquivos originais e duplicatas:** preservados integralmente na pasta [`originais/`](originais/)',
        '- **Formato:** texto plano com cabeçalho `NUM_CANDIDATOS`, `NUM_ESCOLAS`, blocos `# CANDIDATOS`, `# ESCOLAS` e `# MATRIZ_DISTANCIAS`',
        '- **Rastreabilidade:** mapeamento completo disponível em [`manifesto.csv`](manifesto.csv)',
        '',
        '## Tabela de instâncias',
        '',
        '| Arquivo | Origem / Grupo | n | m | Capacidade | Utilização | Duplicatas / Nomes alternativos |',
        '|---|---|---:|---:|---:|---:|---|'
    ]
    for r in manifest_rows:
        md_lines.append(f"| `{r['arquivo']}` | {r['grupo_descricao']} | {r['n_candidatos']} | {r['m_escolas']} | {r['capacidade_total']} | {r['utilizacao_pct']} | {r['arquivos_duplicados']} |")
    
    md_lines.extend([
        '',
        '## Como executar',
        '',
        'Para executar uma instância específica:',
        '```sh',
        './alocacao --instancia instances/instancia_01_prof_A.txt --tipos-por-escola 1 --metodo vnd --runs 1',
        '```',
        '',
        'Para executar todas as 41 instâncias em lote:',
        '```sh',
        'python3 scripts/executar_todas.py',
        '# Ou via Makefile:',
        'make run-all',
        '```',
        ''
    ])
    
    readme_path = INSTANCES_DIR / 'README.md'
    readme_path.write_text('\n'.join(md_lines), encoding='utf-8')
    print(f'   Salvo: {readme_path.relative_to(ROOT)}')
    print('Padronização concluída com sucesso!')

if __name__ == '__main__':
    main()
